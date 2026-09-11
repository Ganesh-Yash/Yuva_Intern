"""Shared helpers for building the internship Word reports."""
from __future__ import annotations

from pathlib import Path

from docx import Document
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_LINE_SPACING
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Inches, Pt, RGBColor


NAVY = RGBColor(0x1F, 0x4E, 0x79)
ACCENT = RGBColor(0x2E, 0x75, 0xB6)
DARK = RGBColor(0x2D, 0x2D, 0x2D)
GRAY = RGBColor(0x55, 0x55, 0x55)
WHITE = RGBColor(0xFF, 0xFF, 0xFF)
HEADER_BG = "1F4E79"
ALT_ROW = "EAF1F8"
CODE_BG = "F4F6F8"


def _set_run_font(run, name="Calibri", size=11, bold=False, italic=False, color=DARK):
    run.font.name = name
    run._element.rPr.rFonts.set(qn("w:eastAsia"), name)
    run.font.size = Pt(size)
    run.bold = bold
    run.italic = italic
    run.font.color.rgb = color


def _shade_cell(cell, hex_color):
    tc = cell._tePr if hasattr(cell, "_tePr") else cell._tc
    tcPr = tc.get_or_add_tcPr()
    shd = OxmlElement("w:shd")
    shd.set(qn("w:fill"), hex_color)
    shd.set(qn("w:val"), "clear")
    tcPr.append(shd)


def _set_cell_borders(cell, color="BFBFBF"):
    tc = cell._tc
    tcPr = tc.get_or_add_tcPr()
    tcBorders = OxmlElement("w:tcBorders")
    for edge in ("top", "left", "bottom", "right"):
        el = OxmlElement(f"w:{edge}")
        el.set(qn("w:val"), "single")
        el.set(qn("w:sz"), "4")
        el.set(qn("w:space"), "0")
        el.set(qn("w:color"), color)
        tcBorders.append(el)
    tcPr.append(tcBorders)


def _set_paragraph_spacing(p, before=0, after=8, line=1.15):
    pf = p.paragraph_format
    pf.space_before = Pt(before)
    pf.space_after = Pt(after)
    pf.line_spacing = line


def add_horizontal_line(paragraph):
    p = paragraph._p
    pPr = p.get_or_add_pPr()
    pBdr = OxmlElement("w:pBdr")
    bottom = OxmlElement("w:bottom")
    bottom.set(qn("w:val"), "single")
    bottom.set(qn("w:sz"), "12")
    bottom.set(qn("w:space"), "1")
    bottom.set(qn("w:color"), "1F4E79")
    pBdr.append(bottom)
    pPr.append(pBdr)


def add_page_number(paragraph):
    run = paragraph.add_run()
    fld_begin = OxmlElement("w:fldChar")
    fld_begin.set(qn("w:fldCharType"), "begin")
    instr = OxmlElement("w:instrText")
    instr.set(qn("xml:space"), "preserve")
    instr.text = " PAGE "
    fld_end = OxmlElement("w:fldChar")
    fld_end.set(qn("w:fldCharType"), "end")
    run._r.append(fld_begin)
    run._r.append(instr)
    run._r.append(fld_end)


class InternshipReport:
    def __init__(self, title, subtitle, author, date, task, tools):
        self.doc = Document()
        self.fig_count = 0
        self.table_count = 0
        section = self.doc.sections[0]
        section.top_margin = Cm(2.0)
        section.bottom_margin = Cm(2.0)
        section.left_margin = Cm(2.2)
        section.right_margin = Cm(2.2)
        section.page_width = Cm(21.0)
        section.page_height = Cm(29.7)

        header = section.header
        header.is_linked_to_previous = False
        hp = header.paragraphs[0]
        hp.alignment = WD_ALIGN_PARAGRAPH.RIGHT
        run = hp.add_run(subtitle)
        _set_run_font(run, size=9, italic=True, color=ACCENT)

        footer = section.footer
        footer.is_linked_to_previous = False
        fp = footer.paragraphs[0]
        fp.alignment = WD_ALIGN_PARAGRAPH.CENTER
        r1 = fp.add_run(f"{author}  ·  {title}  ·  Page ")
        _set_run_font(r1, size=8, color=GRAY)
        add_page_number(fp)
        r2 = fp.add_run("")
        _set_run_font(r2, size=8, color=GRAY)

        styles = self.doc.styles
        styles["Normal"].font.name = "Calibri"
        styles["Normal"].font.size = Pt(11)
        styles["Normal"].font.color.rgb = DARK

        self._cover(title, subtitle, author, date, task, tools)

    def _cover(self, title, subtitle, author, date, task, tools):
        p = self.doc.add_paragraph()
        p.alignment = WD_ALIGN_PARAGRAPH.LEFT
        run = p.add_run(title)
        _set_run_font(run, size=26, bold=True, color=NAVY)
        _set_paragraph_spacing(p, after=2)

        p = self.doc.add_paragraph()
        run = p.add_run(subtitle)
        _set_run_font(run, size=14, italic=True, color=ACCENT)
        add_horizontal_line(p)
        _set_paragraph_spacing(p, after=16)

        meta = [
            ("Submitted by", author),
            ("Date", date),
            ("Task", task),
            ("Tools Used", tools),
        ]
        table = self.doc.add_table(rows=len(meta), cols=2)
        table.autofit = True
        for i, (k, v) in enumerate(meta):
            c0, c1 = table.rows[i].cells
            c0.text = ""
            c1.text = ""
            p0 = c0.paragraphs[0]
            r0 = p0.add_run(k)
            _set_run_font(r0, size=11, bold=True, color=NAVY)
            p1 = c1.paragraphs[0]
            r1 = p1.add_run(v)
            _set_run_font(r1, size=11, color=DARK)
            _shade_cell(c0, "EAF1F8")
            _set_cell_borders(c0, "D0D7DE")
            _set_cell_borders(c1, "D0D7DE")
        self.doc.add_paragraph()

    def heading(self, text, level=1):
        p = self.doc.add_paragraph()
        size = 16 if level == 1 else 13
        run = p.add_run(text)
        _set_run_font(run, size=size, bold=True, color=NAVY)
        if level == 1:
            add_horizontal_line(p)
        _set_paragraph_spacing(p, before=14 if level == 1 else 10, after=8)
        return p

    def para(self, text):
        p = self.doc.add_paragraph()
        run = p.add_run(text)
        _set_run_font(run, size=11, color=DARK)
        _set_paragraph_spacing(p, after=8, line=1.15)
        p.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
        return p

    def bullet(self, text):
        p = self.doc.add_paragraph(style="List Bullet")
        p.clear()
        run = p.add_run(text)
        _set_run_font(run, size=11, color=DARK)
        _set_paragraph_spacing(p, after=4, line=1.15)
        return p

    def code(self, text):
        p = self.doc.add_paragraph()
        run = p.add_run(text)
        _set_run_font(run, name="Consolas", size=9, color=RGBColor(0x1E, 0x1E, 0x1E))
        _set_paragraph_spacing(p, before=4, after=8, line=1.08)
        pPr = p._p.get_or_add_pPr()
        shd = OxmlElement("w:shd")
        shd.set(qn("w:fill"), CODE_BG)
        shd.set(qn("w:val"), "clear")
        pPr.append(shd)
        return p

    def table(self, headers, rows, caption=None):
        self.table_count += 1
        if caption:
            cap = self.doc.add_paragraph()
            run = cap.add_run(f"Table {self.table_count}. {caption}")
            _set_run_font(run, size=10, italic=True, color=GRAY)
            _set_paragraph_spacing(cap, before=8, after=4)

        table = self.doc.add_table(rows=1 + len(rows), cols=len(headers))
        table.alignment = WD_TABLE_ALIGNMENT.CENTER
        table.autofit = True
        for j, h in enumerate(headers):
            cell = table.rows[0].cells[j]
            cell.text = ""
            p = cell.paragraphs[0]
            p.alignment = WD_ALIGN_PARAGRAPH.CENTER
            run = p.add_run(str(h))
            _set_run_font(run, size=10, bold=True, color=WHITE)
            _shade_cell(cell, HEADER_BG)
            _set_cell_borders(cell, "1F4E79")
        for i, row in enumerate(rows):
            for j, val in enumerate(row):
                cell = table.rows[i + 1].cells[j]
                cell.text = ""
                p = cell.paragraphs[0]
                p.alignment = WD_ALIGN_PARAGRAPH.CENTER if j else WD_ALIGN_PARAGRAPH.LEFT
                run = p.add_run(str(val))
                _set_run_font(run, size=10, color=DARK)
                if i % 2 == 1:
                    _shade_cell(cell, ALT_ROW)
                _set_cell_borders(cell, "D0D7DE")
        self.doc.add_paragraph()
        return table

    def image(self, path, caption, width=6.2):
        self.fig_count += 1
        p = self.doc.add_paragraph()
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p.add_run().add_picture(str(path), width=Inches(width))
        _set_paragraph_spacing(p, before=6, after=2)
        cap = self.doc.add_paragraph()
        cap.alignment = WD_ALIGN_PARAGRAPH.CENTER
        run = cap.add_run(f"Figure {self.fig_count}. {caption}")
        _set_run_font(run, size=10, italic=True, color=GRAY)
        _set_paragraph_spacing(cap, before=0, after=10)

    def save(self, path):
        Path(path).parent.mkdir(parents=True, exist_ok=True)
        self.doc.save(str(path))
        return str(path)
